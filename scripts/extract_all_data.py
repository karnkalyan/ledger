import os
import glob
import re
import json
import fitz
import openpyxl
from datetime import datetime

def clean_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(',', '').replace(' ', '')
    if not s or s == '-':
        return 0.0
    try:
        return float(s)
    except:
        return 0.0

def parse_date_str(d):
    if not d:
        return None
    if isinstance(d, datetime):
        return d.strftime('%Y-%m-%d')
    s = str(d).strip()
    # YYYY-MM-DD or YYYY-MM-DD HH:MM:SS
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    # DD.MM.YY or DD.MM.YYYY
    m = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{2,4})', s)
    if m:
        day = int(m.group(1))
        mon = int(m.group(2))
        yr = int(m.group(3))
        if yr < 100:
            yr = 2000 + yr
        if yr >= 2070:
            # BS date
            return f"BS:{yr}-{mon:02d}-{day:02d}"
        return f"{yr:04d}-{mon:02d}-{day:02d}"
    # YYYY.MM.DD
    m = re.match(r'^(\d{4})\.(\d{1,2})\.(\d{1,2})', s)
    if m:
        yr = int(m.group(1))
        mon = int(m.group(2))
        day = int(m.group(3))
        if yr >= 2070:
            return f"BS:{yr}-{mon:02d}-{day:02d}"
        return f"{yr:04d}-{mon:02d}-{day:02d}"
    return s

def extract_pdf_statements():
    pdf_files = sorted(glob.glob(r'companyDetails/statement-*.pdf'))
    # deduplicate by base name without " (1)"
    unique_pdfs = []
    seen_bases = set()
    for p in pdf_files:
        base = os.path.basename(p).replace(' (1)', '')
        if base in seen_bases:
            continue
        seen_bases.add(base)
        unique_pdfs.append(p)

    all_txs = []
    seen_keys = set()

    for pdf in unique_pdfs:
        doc = fitz.open(pdf)
        for page in doc:
            lines = [l.strip() for l in page.get_text('text').split('\n') if l.strip()]
            start_idx = 0
            end_idx = len(lines)
            for i, l in enumerate(lines):
                if l.startswith('It is understood and agreed'):
                    end_idx = i
                    break
            for i in range(end_idx):
                if lines[i] in ['Opening Balance', 'Closing Balance'] or re.match(r'^\d{4}-\d{2}-\d{2}', lines[i]):
                    if i > 0 and lines[i-1] in ['From :', 'To :']:
                        continue
                    start_idx = i
                    if lines[i].count('-') == 2 and i + 1 < end_idx and lines[i+1] == 'Opening Balance':
                        start_idx = i
                        break
                    elif re.match(r'^\d{4}-\d{2}-\d{2}', lines[i]):
                        start_idx = i
                        break

            table_lines = lines[start_idx:end_idx]
            i = 0
            while i < len(table_lines):
                line = table_lines[i]
                if re.match(r'^\d{4}-\d{2}-\d{2}', line):
                    tx_date = line[:10]
                    desc_parts = []
                    j = i + 1
                    while j < len(table_lines) and not re.match(r'^\d{4}-\d{2}-\d{2}', table_lines[j]):
                        desc_parts.append(table_lines[j])
                        j += 1

                    if len(desc_parts) >= 3:
                        sn = None
                        if desc_parts[-1].isdigit():
                            sn = desc_parts[-1]
                            desc_parts_no_sn = desc_parts[:-1]
                        else:
                            desc_parts_no_sn = desc_parts

                        if len(desc_parts_no_sn) >= 3:
                            bal_str = desc_parts_no_sn[-1]
                            dep_str = desc_parts_no_sn[-2]
                            wd_str = desc_parts_no_sn[-3]
                            desc = ' '.join(desc_parts_no_sn[:-3]).strip()
                            if not desc and desc_parts_no_sn:
                                desc = 'Bank Transaction'

                            wd_amt = clean_num(wd_str)
                            dep_amt = clean_num(dep_str)
                            bal_amt = clean_num(bal_str)

                            # Skip pure opening/closing balance rows as transactions
                            if desc.lower() not in ['opening balance', 'closing balance']:
                                tx_type = 'DEPOSIT' if dep_amt > 0 else 'WITHDRAWAL'
                                amt = dep_amt if dep_amt > 0 else wd_amt
                                key = (tx_date, tx_type, round(amt, 2), round(bal_amt, 2), desc[:40])
                                if key not in seen_keys and amt > 0:
                                    seen_keys.add(key)
                                    all_txs.append({
                                        'date': tx_date,
                                        'type': tx_type,
                                        'amount': amt,
                                        'balance': bal_amt,
                                        'description': desc,
                                        'reference': f"SN-{sn}" if sn else None,
                                        'source_pdf': os.path.basename(pdf),
                                    })
                    i = j
                else:
                    i += 1

    return all_txs

def extract_everest_bank():
    path = r'companyDetails/Desktop/BANK EVEREST AV.xlsx'
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    txs = []
    
    # 1. Sheet current -> Account 04400105200931
    if 'current' in wb.sheetnames:
        ws = wb['current']
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[3:]:
            if not any(r): continue
            sno, mitti, dt, particular, chq, remarks, amt_dr = (r + (None,)*7)[:7]
            amt_cr = r[7] if len(r) > 7 else None
            bal = r[8] if len(r) > 8 else None
            if not particular and not amt_dr and not amt_cr: continue
            dr = clean_num(amt_dr)
            cr = clean_num(amt_cr)
            if dr == 0 and cr == 0: continue
            tx_type = 'DEPOSIT' if cr > 0 else 'WITHDRAWAL'
            amt = cr if cr > 0 else dr
            d_str = parse_date_str(dt) or parse_date_str(mitti) or '2023-05-15'
            txs.append({
                'accountNumber': '04400105200931',
                'accountName': 'Everest Bank Current A/C',
                'date': d_str,
                'type': tx_type,
                'amount': amt,
                'balance': clean_num(bal) if bal else None,
                'description': str(particular or '').strip() or 'Everest Bank Transaction',
                'chequeNumber': str(chq).strip() if chq else None,
                'reference': str(remarks).strip() if remarks else None,
            })

    # 2. Sheet od -> Account 04405501200137
    if 'od' in wb.sheetnames:
        ws = wb['od']
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[3:]:
            if not any(r): continue
            sno, mitti, dt, particular, chq, remarks, amt_dr = (r + (None,)*7)[:7]
            amt_cr = r[7] if len(r) > 7 else None
            bal = r[8] if len(r) > 8 else None
            if not particular and not amt_dr and not amt_cr: continue
            dr = clean_num(amt_dr)
            cr = clean_num(amt_cr)
            if dr == 0 and cr == 0: continue
            tx_type = 'DEPOSIT' if cr > 0 else 'WITHDRAWAL'
            amt = cr if cr > 0 else dr
            d_str = parse_date_str(dt) or parse_date_str(mitti) or '2023-06-01'
            txs.append({
                'accountNumber': '04405501200137',
                'accountName': 'Everest Bank Overdraft (OD) A/C',
                'date': d_str,
                'type': tx_type,
                'amount': amt,
                'balance': clean_num(bal) if bal else None,
                'description': str(particular or '').strip() or 'Everest Bank OD Transaction',
                'chequeNumber': str(chq).strip() if chq else None,
                'reference': str(remarks).strip() if remarks else None,
            })

    return txs

def extract_products():
    path = r'companyDetails/SIMUL STOCK.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Sheet1']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    seen_sku = set()

    for idx, r in enumerate(rows[5:], 1):
        if not any(r): continue
        sno = r[0]
        name = r[1]
        if not name or str(name).strip().lower() in ['total', 'grand total', 's.no', 'product name']:
            continue
        unit = str(r[2] or 'pcs').strip()
        code = str(r[3] or '').strip()
        rack = str(r[4] or '').strip()
        qty = clean_num(r[5])
        vat_rate = clean_num(r[6])
        cat = str(r[7] or 'Broadcasting Hardware').strip()
        landing_cost = clean_num(r[9]) if len(r) > 9 else 0.0
        selling_price = clean_num(r[11]) if len(r) > 11 else (vat_rate if vat_rate > 0 else landing_cost * 1.25)
        if selling_price == 0:
            selling_price = landing_cost * 1.25 if landing_cost > 0 else 1000.0

        p_name = str(name).strip()
        sku = code if code else f"SIM-PRD-{idx:03d}"
        if sku in seen_sku:
            sku = f"{sku}-{idx}"
        seen_sku.add(sku)

        products.append({
            'sku': sku,
            'name': p_name,
            'unit': unit or 'pcs',
            'category': cat,
            'stockQty': qty,
            'purchasePrice': landing_cost if landing_cost > 0 else round(selling_price * 0.75, 2),
            'salePrice': selling_price if selling_price > 0 else round(landing_cost * 1.25, 2),
            'reorderLevel': 2,
        })
    return products

def extract_purchases():
    path = r'companyDetails/SIMUL PURCHASE.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True)
    orders = []
    
    for sheet_name in ['79-80', '80-81']:
        if sheet_name not in wb.sheetnames: continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        current_vendor = None
        current_date = None
        current_inv = None

        for idx, r in enumerate(rows[5:], 1):
            if not any(r): continue
            sno, nep_d, eng_d, vendor, invoice, prod, qty, rate, total, vat = (r + (None,)*10)[:10]
            if vendor and str(vendor).strip():
                current_vendor = str(vendor).strip()
            if eng_d or nep_d:
                current_date = parse_date_str(eng_d) or parse_date_str(nep_d)
            if invoice and str(invoice).strip():
                current_inv = str(invoice).strip()
            
            p_name = str(prod or '').strip()
            tot = clean_num(total)
            q = clean_num(qty)
            rt = clean_num(rate)
            if not p_name and tot == 0: continue
            if q == 0 and rt > 0 and tot > 0:
                q = 1.0
            if rt == 0 and tot > 0 and q > 0:
                rt = tot / q

            orders.append({
                'fiscalYear': sheet_name,
                'vendor': current_vendor or 'Local Vendor',
                'invoiceNumber': current_inv or f"PUR-{sheet_name}-{idx:03d}",
                'date': current_date or '2023-04-01',
                'productName': p_name or 'Procurement Item',
                'quantity': max(1.0, q),
                'unitPrice': rt if rt > 0 else tot,
                'total': tot if tot > 0 else q * rt,
                'vat': clean_num(vat),
            })
    return orders

def extract_sales():
    path = r'companyDetails/SIMUL SALES.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True)
    sales = []
    payments = []

    for sheet_name in ['SALES 79-80', '80-81']:
        if sheet_name not in wb.sheetnames: continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        current_client = None
        current_date = None
        current_bill = None

        start_row = 6
        for idx, r in enumerate(rows[start_row:], 1):
            if not any(r): continue
            # Check structure:
            if sheet_name == 'SALES 79-80':
                sno, mitti, bill_no, client, prod, qty, rate, total, vat = (r + (None,)*9)[:9]
                dt = parse_date_str(mitti)
            else:
                sno, mitti, eng_date, bill_no, client, prod, qty, rate, total, vat = (r + (None,)*10)[:10]
                dt = parse_date_str(eng_date) or parse_date_str(mitti)

            if client and str(client).strip():
                current_client = str(client).strip()
            if dt:
                current_date = dt
            if bill_no and str(bill_no).strip():
                current_bill = str(bill_no).strip()

            p_name = str(prod or '').strip()
            tot = clean_num(total)
            q = clean_num(qty)
            rt = clean_num(rate)
            if not p_name and tot == 0: continue
            if q == 0 and tot > 0: q = 1.0
            if rt == 0 and tot > 0 and q > 0: rt = tot / q

            sales.append({
                'fiscalYear': sheet_name,
                'client': current_client or 'Direct Client',
                'billNumber': current_bill or f"BILL-{sheet_name}-{idx:03d}",
                'date': current_date or '2023-07-01',
                'productName': p_name or 'Broadcasting Equipment',
                'quantity': max(1.0, q),
                'unitPrice': rt if rt > 0 else tot,
                'total': tot if tot > 0 else q * rt,
                'vat': clean_num(vat),
            })

    if 'PAYMENT ' in wb.sheetnames or 'PAYMENT' in wb.sheetnames:
        p_sheet = 'PAYMENT ' if 'PAYMENT ' in wb.sheetnames else 'PAYMENT'
        ws = wb[p_sheet]
        rows = list(ws.iter_rows(values_only=True))
        for idx, r in enumerate(rows[5:], 1):
            if not any(r): continue
            sno, nep_d, eng_d, party, amt, bill_no, remarks, mode = (r + (None,)*8)[:8]
            party_name = str(party or '').strip()
            a = clean_num(amt)
            if not party_name and a == 0: continue
            dt = parse_date_str(eng_d) or parse_date_str(nep_d) or '2023-08-01'
            payments.append({
                'client': party_name,
                'amount': a,
                'date': dt,
                'billNumber': str(bill_no).strip() if bill_no else None,
                'method': str(mode or 'Bank Transfer').strip(),
                'remarks': str(remarks).strip() if remarks else None,
            })

    return sales, payments

def extract_office_expenses():
    path = r'companyDetails/SIMULCAST STATEMENT.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True)
    expenses = []
    if 'Sheet2' in wb.sheetnames:
        ws = wb['Sheet2']
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[4:]:
            if not any(r): continue
            sno, d1, d2, particular, a1, a2 = (r + (None,)*6)[:6]
            amt = clean_num(a1) if clean_num(a1) > 0 else clean_num(a2)
            part = str(particular or '').strip()
            if not part and amt == 0: continue
            d_str = parse_date_str(d1) or parse_date_str(d2) or '2023-04-01'
            expenses.append({
                'date': d_str,
                'particular': part or 'General Office Expense',
                'amount': amt,
            })
    return expenses

def extract_contacts(purchases, sales, payments):
    clients = set()
    vendors = set()

    for p in purchases:
        v = p.get('vendor')
        if v and len(v) > 2 and v.lower() not in ['s no', 'vendor name', 'total']:
            vendors.add(v)

    for s in sales:
        c = s.get('client')
        if c and len(c) > 2 and c.lower() not in ['s no', 'client name', 'total']:
            clients.add(c)

    for pay in payments:
        c = pay.get('client')
        if c and len(c) > 2:
            clients.add(c)

    # Also check Desktop client ledgers
    cl_path = r'companyDetails/Desktop/client ledger/CLIENT LEDGER.xlsx'
    if os.path.exists(cl_path):
        wb = openpyxl.load_workbook(cl_path, data_only=True)
        for s in wb.sheetnames:
            if 'sheet' not in s.lower():
                clients.add(s.strip())

    # Check PARTY LEDGER
    pl_path = r'companyDetails/Desktop/PARTY LEDGER.xlsx'
    if os.path.exists(pl_path):
        wb = openpyxl.load_workbook(pl_path, data_only=True)
        for s in wb.sheetnames:
            if 'sheet' not in s.lower():
                vendors.add(s.strip())

    return sorted(list(clients)), sorted(list(vendors))

def main():
    print("Extracting Siddhartha Bank PDFs...")
    pdf_txs = extract_pdf_statements()
    print(f"-> Extracted {len(pdf_txs)} Siddhartha Bank transactions.")

    print("Extracting Everest Bank Excel...")
    eb_txs = extract_everest_bank()
    print(f"-> Extracted {len(eb_txs)} Everest Bank transactions.")

    print("Extracting Products & Stock...")
    products = extract_products()
    print(f"-> Extracted {len(products)} products.")

    print("Extracting Purchases...")
    purchases = extract_purchases()
    print(f"-> Extracted {len(purchases)} purchase line items.")

    print("Extracting Sales & Payments...")
    sales, payments = extract_sales()
    print(f"-> Extracted {len(sales)} sales items, {len(payments)} payments.")

    print("Extracting Office Expenses...")
    office_expenses = extract_office_expenses()
    print(f"-> Extracted {len(office_expenses)} office expenses.")

    clients, vendors = extract_contacts(purchases, sales, payments)
    print(f"-> Found {len(clients)} unique clients, {len(vendors)} unique vendors.")

    data = {
        'siddharthaBank': {
            'bankName': 'Siddhartha Bank Ltd',
            'accountName': 'Simulcast Technologies Pvt Ltd',
            'accountNumber': '02017249683',
            'accountType': 'CANR (Current Account)',
            'currency': 'NPR',
            'openingBalance': 5795.00,
            'transactions': pdf_txs,
        },
        'products': products,
        'purchases': purchases,
        'sales': sales,
        'payments': payments,
        'officeExpenses': office_expenses,
        'clients': clients,
        'vendors': vendors,
    }

    out_file = r'scripts/extracted_company_data.json'
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"SUCCESS: Saved extracted data to {out_file}")

if __name__ == '__main__':
    main()
